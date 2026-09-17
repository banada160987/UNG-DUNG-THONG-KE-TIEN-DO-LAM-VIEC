import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_plan_docx():
    doc = docx.Document()

    # Căn lề chuẩn Nghị định 30/2020/NĐ-CP:
    # Lề trên: 2.0 cm, Lề dưới: 2.0 cm, Lề trái: 3.0 cm, Lề phải: 1.5 cm
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)

    # Style Normal mặc định
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # ==================== PHẦN HEADER ĐẦU TRANG ====================
    # Dùng bảng 1 hàng 2 cột không viền để chia Quốc hiệu & Tên cơ quan
    header_tbl = doc.add_table(rows=1, cols=2)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False

    # Cột 1: Cơ quan ban hành (Rộng ~7.5cm)
    cell_left = header_tbl.cell(0, 0)
    cell_left.width = Cm(7.5)
    set_cell_margins(cell_left, top=0, bottom=0, left=0, right=100)

    p_cq1 = cell_left.paragraphs[0]
    p_cq1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cq1.paragraph_format.space_before = Pt(0)
    p_cq1.paragraph_format.space_after = Pt(2)
    p_cq1.paragraph_format.line_spacing = 1.0
    r_cq1 = p_cq1.add_run("SỞ GIÁO DỤC VÀ ĐÀO TẠO ĐẮK LẮK")
    r_cq1.font.name = 'Times New Roman'
    r_cq1.font.size = Pt(12)
    r_cq1.font.color.rgb = RGBColor(0, 0, 0)

    p_cq2 = cell_left.add_paragraph()
    p_cq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cq2.paragraph_format.space_before = Pt(0)
    p_cq2.paragraph_format.space_after = Pt(2)
    p_cq2.paragraph_format.line_spacing = 1.0
    r_cq2 = p_cq2.add_run("TRƯỜNG THPT CAO BÁ QUÁT")
    r_cq2.font.name = 'Times New Roman'
    r_cq2.font.size = Pt(12.5)
    r_cq2.bold = True
    r_cq2.font.color.rgb = RGBColor(0, 0, 0)

    p_line_cq = cell_left.add_paragraph()
    p_line_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line_cq.paragraph_format.space_before = Pt(0)
    p_line_cq.paragraph_format.space_after = Pt(4)
    r_line_cq = p_line_cq.add_run("————————")
    r_line_cq.bold = True
    r_line_cq.font.size = Pt(10)

    p_so = cell_left.add_paragraph()
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.space_before = Pt(0)
    p_so.paragraph_format.space_after = Pt(0)
    r_so = p_so.add_run("Số:       /KH-TrTHPTCBQ")
    r_so.font.name = 'Times New Roman'
    r_so.font.size = Pt(12.5)

    # Cột 2: Quốc hiệu Tiêu ngữ (Rộng ~9.0cm)
    cell_right = header_tbl.cell(0, 1)
    cell_right.width = Cm(9.0)
    set_cell_margins(cell_right, top=0, bottom=0, left=100, right=0)

    p_qh1 = cell_right.paragraphs[0]
    p_qh1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_qh1.paragraph_format.space_before = Pt(0)
    p_qh1.paragraph_format.space_after = Pt(2)
    p_qh1.paragraph_format.line_spacing = 1.0
    r_qh1 = p_qh1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r_qh1.font.name = 'Times New Roman'
    r_qh1.font.size = Pt(12)
    r_qh1.bold = True
    r_qh1.font.color.rgb = RGBColor(0, 0, 0)

    p_qh2 = cell_right.add_paragraph()
    p_qh2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_qh2.paragraph_format.space_before = Pt(0)
    p_qh2.paragraph_format.space_after = Pt(2)
    p_qh2.paragraph_format.line_spacing = 1.0
    r_qh2 = p_qh2.add_run("Độc lập - Tự do - Hạnh phúc")
    r_qh2.font.name = 'Times New Roman'
    r_qh2.font.size = Pt(13)
    r_qh2.bold = True
    r_qh2.font.color.rgb = RGBColor(0, 0, 0)

    p_line_qh = cell_right.add_paragraph()
    p_line_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line_qh.paragraph_format.space_before = Pt(0)
    p_line_qh.paragraph_format.space_after = Pt(4)
    r_line_qh = p_line_qh.add_run("—————————————")
    r_line_qh.bold = True
    r_line_qh.font.size = Pt(10)

    p_date = cell_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(0)
    r_date = p_date.add_run("Tân An, ngày 05 tháng 9 năm 2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.italic = True

    # ==================== TIÊU ĐỀ KẾ HOẠCH ====================
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(3)
    r_t1 = p_title.add_run("KẾ HOẠCH")
    r_t1.font.name = 'Times New Roman'
    r_t1.font.size = Pt(14)
    r_t1.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    r_t2 = p_sub.add_run("Công tác trọng tâm tháng 9 năm 2026")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(13)
    r_t2.bold = True

    p_line_title = doc.add_paragraph()
    p_line_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line_title.paragraph_format.space_before = Pt(0)
    p_line_title.paragraph_format.space_after = Pt(14)
    r_lt = p_line_title.add_run("————————")
    r_lt.font.size = Pt(10)
    r_lt.bold = True

    # Helpers
    def add_sec_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.bold = True
        return p

    def add_sub_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.bold = True
        return p

    def add_p(text, indent=True, bold_prefix=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.18
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.0)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(13)
        rc = p.add_run(text)
        rc.font.name = 'Times New Roman'
        rc.font.size = Pt(13)
        return p

    # ==================== PHẦN A ====================
    add_sec_h("A. ĐÁNH GIÁ CÔNG TÁC THÁNG 8/2026 (CÔNG TÁC CHUẨN BỊ NĂM HỌC MỚI)")
    add_p("Trong tháng 8/2026, Nhà trường đã tập trung chỉ đạo và tổ chức thực hiện toàn diện các nhiệm vụ chuẩn bị cho năm học mới 2026 – 2027 và tổ chức thành công các sự kiện lớn của đơn vị, kết quả cụ thể như sau:")

    add_p("100% cán bộ, giáo viên, nhân viên (CB-GV-NV) tham gia nghiêm túc đợt Bồi dưỡng chính trị hè 2026 do Sở Giáo dục và Đào tạo phối hợp tổ chức và các lớp tập huấn chuyên môn Chương trình GDPT 2018; nộp bài thu hoạch đầy đủ, đảm bảo chất lượng theo quy định.", True, "1. Công tác bồi dưỡng chính trị và chuyên môn hè: ")

    add_p("Hoàn thành rà soát, sắp xếp ổn định biên chế học sinh các khối lớp 10, 11, 12; phân công Giáo viên chủ nhiệm (GVCN) và Giáo viên bộ môn (GVBM) đúng năng lực, trình độ chuyên môn, đảm bảo cân đối về định mức giờ dạy theo quy định.", True, "2. Công tác tổ chức đội ngũ và biên chế lớp học:\n- ")
    add_p("Thực hiện quy trình kiện toàn cơ cấu tổ chức bộ máy nhà trường năm học 2026 – 2027: Thành lập 07 Tổ Chuyên môn và 01 Tổ Văn phòng (giảm 02 tổ chuyên môn so với năm học trước nhằm tinh gọn, nâng cao hiệu quả); bổ nhiệm mới Thư ký hội đồng, 02 Tổ phó; bổ nhiệm lại 07 Tổ trưởng và 02 Tổ phó chuyên môn đúng quy định.", True, "- ")

    add_p("100% Tổ chuyên môn đã hoàn thành việc xây dựng Kế hoạch giáo dục của tổ (phân phối chương trình, ma trận kiểm tra đánh giá định kỳ, kế hoạch dạy học 2 buổi/ngày) bám sát yêu cầu cần đạt của Chương trình GDPT 2018.", True, "3. Công tác xây dựng kế hoạch giáo dục: ")

    add_p("Hoàn thành kiểm tra, tu sửa, bổ sung bàn ghế, thiết bị phòng học bộ môn, vệ sinh cảnh quan môi trường sư phạm xanh – sạch – đẹp; chuẩn bị đầy đủ hạ tầng CNTT, phòng máy tính, hệ thống phần mềm và kho học liệu số sẵn sàng ứng dụng công nghệ và Trí tuệ nhân tạo (AI) vào giảng dạy ngay từ đầu năm học.", True, "4. Công tác cơ sở vật chất, học liệu và ứng dụng công nghệ: ")

    add_p("Tổ chức thành công tốt đẹp Lễ Kỷ niệm 30 năm Ngày thành lập trường (1996 – 2026). Sự kiện đã lan tỏa truyền thống vẻ vang của nhà trường, tạo động lực tinh thần to lớn đối với toàn thể cán bộ, giáo viên, nhân viên và học sinh; nhận được sự đánh giá cao từ Lãnh đạo các cấp, các thế hệ cựu học sinh và nhân dân địa phương.", True, "5. Công tác tổ chức sự kiện: ")

    # ==================== PHẦN B ====================
    add_sec_h("B. KẾ HOẠCH CÔNG TÁC TRỌNG TÂM THÁNG 9/2026")

    add_sub_h("I. CÔNG TÁC TƯ TƯỞNG, CHÍNH TRỊ VÀ THI ĐUA ĐẦU NĂM HỌC")
    add_p("1. Tuyên truyền không khí thi đua sôi nổi nhân Ngày toàn dân đưa trẻ đến trường (05/9); định hướng tư tưởng, tác phong làm việc nghiêm túc, sẵn sàng tâm thế bước vào năm học mới 2026 – 2027 cho 100% CB-GV-NV và học sinh.")
    add_p("2. Phát động phong trào thi đua \"Dạy tốt – Học tốt\", xây dựng \"Trường học hạnh phúc, kỷ cương, an toàn và thân thiện\" ngay từ những tuần học đầu tiên.")
    add_p("3. Triển khai đầy đủ, kịp thời các văn bản chỉ đạo, hướng dẫn thực hiện nhiệm vụ năm học 2026 – 2027 của Bộ GD&ĐT, UBND tỉnh và Sở GD&ĐT Đắk Lắk.")

    add_sub_h("II. CÔNG TÁC CHUYÊN MÔN VÀ DẠY HỌC (NHIỆM VỤ TRỌNG TÂM)")
    add_p("Tổ chức Lễ Khai giảng năm học 2026 – 2027 vào sáng ngày 05/9/2026 đảm bảo trang trọng, ý nghĩa, ngắn gọn và an toàn tuyệt đối; tổ chức thành công Hội nghị Cán bộ, Viên chức, Người lao động năm học 2026 – 2027 đúng quy trình, dân chủ; hoàn tất báo cáo tổng kết Lễ kỷ niệm 30 năm thành lập trường gửi về Sở GD&ĐT.", True, "1. Tổ chức các sự kiện đầu năm học: ")

    add_p("Thực hiện nghiêm túc khung kế hoạch thời gian năm học 2026 – 2027; triển khai đồng bộ, hiệu quả Chương trình GDPT 2018 cho tất cả các khối lớp 10, 11, 12; tổ chức dạy học 2 buổi/ngày đảm bảo chất lượng, đúng phân phối chương trình và đáp ứng nhu cầu phân hóa năng lực học sinh; ổn định thời khóa biểu, siết chặt kỷ cương chuyên môn và quy định đạo đức nhà giáo ngay từ tiết học đầu tiên.", True, "2. Thực hiện kế hoạch dạy và học: ")

    add_p("Tổ chức kiểm tra, đánh giá phân loại chất lượng học sinh đầu năm ở các môn học trọng điểm để có kế hoạch phụ đạo học sinh yếu và bồi dưỡng học sinh khá giỏi; thành lập, phân công giáo viên phụ trách và triển khai ôn luyện Đội tuyển dự thi Kỳ thi chọn học sinh giỏi quốc gia năm học 2026 – 2027.", True, "3. Khảo sát chất lượng & Bồi dưỡng học sinh giỏi: ")

    add_p("Đẩy mạnh đổi mới phương pháp giảng dạy, đổi mới kiểm tra đánh giá theo định dạng đề thi mới; đưa hệ thống phần mềm khảo thí thông minh và các Trợ lý AI (soạn giáo án 5512, chấm tự luận, luyện tiếng Anh, thi lập trình DMOJ) vào hỗ trợ công tác chuyên môn; ban hành quyết định thành lập và chính thức đưa vào hoạt động các Câu lạc bộ học thuật, Tin học/AI, Tiếng Anh, Thể thao và Nghệ thuật.", True, "4. Đổi mới phương pháp và chuyển đổi số dạy học: ")

    add_sub_h("III. CÔNG TÁC QUẢN LÝ HỌC SINH, NỀ NẾP VÀ GIÁO DỤC KỸ NĂNG")
    add_p("1. GVCN phối hợp chặt chẽ với Đoàn Thanh niên và Ban Quản lý nề nếp ổn định sĩ số lớp, quán triệt nội quy trường lớp, trang phục, đồng phục của học sinh; áp dụng điểm danh điện tử để kiểm soát chuyên cần hàng ngày và liên lạc kịp thời với gia đình học sinh.")
    add_p("2. Phối hợp chặt chẽ giữa Lãnh đạo trường, GVCN, Đoàn Thanh niên và Ban Đại diện CMHS trong công tác quản lý, giáo dục đạo đức, phòng chống bạo lực học đường; ký cam kết chấp hành nghiêm Luật Giao thông đường bộ; tổ chức Cuộc thi tìm hiểu và phòng chống tác hại của thuốc lá, thuốc lá điện tử trong học đường.")

    add_sub_h("IV. CÔNG TÁC CƠ SỞ VẬT CHẤT, THIẾT BỊ VÀ TÀI CHÍNH")
    add_p("1. Kiểm tra, bảo dưỡng định kỳ hệ thống máy tính phòng học Tin học, máy chiếu, bảng tương tác, đường truyền mạng Internet tại các phòng học và phòng bộ môn.")
    add_p("2. Đảm bảo cung ứng đầy đủ thiết bị, hóa chất thí nghiệm, đồ dùng dạy học phục vụ các tiết thực hành môn KHTN, Tin học, Ngoại ngữ.")
    add_p("3. Rà soát hệ thống điện nước, công tác phòng cháy chữa cháy, đảm bảo an toàn tuyệt đối trong khuôn viên nhà trường mùa mưa bão.")
    add_p("4. Bộ phận Kế toán hoàn thành hồ sơ dự trù kinh phí hoạt động chuyên môn, các hội thi và chi trả chế độ đầu năm học cho CB-GV-NV kịp thời, đúng quy định.")

    add_sub_h("V. CÔNG TÁC ĐOÀN THỂ VÀ PHỐI HỢP NGOẠI KHÓA")
    add_p("1. Chỉ đạo Đoàn Thanh niên xây dựng và ký kết Quy chế phối hợp công tác năm học 2026 – 2027 với Lãnh đạo nhà trường.")
    add_p("2. Hướng dẫn các Chi đoàn tổ chức Đại hội Chi đoàn nhiệm kỳ mới; chuẩn bị tốt các điều kiện để tiến tới Đại hội Đại biểu Đoàn TNCS Hồ Chí Minh trường THPT Cao Bá Quát nhiệm kỳ 2026 – 2027.")

    # ==================== PHẦN C ====================
    add_sec_h("C. TIẾN ĐỘ THỰC HIỆN CÔNG TÁC TRONG THÁNG 9/2026")
    add_p("Tổ chức Lễ Khai giảng năm học 2026 – 2027 (ngày 05/9); ổn định biên chế lớp, ban hành Thời khóa biểu số 1, bắt đầu thực hiện chương trình học kỳ I từ ngày 07/9/2026.", True, "• Tuần 1 (Từ 01/9 – 06/9/2026): ")
    add_p("Thực hiện dạy học chính khóa và dạy học 2 buổi/ngày theo TKB; các Tổ chuyên môn họp duyệt Kế hoạch giáo dục và kế hoạch bài dạy 5512; tổ chức Hội nghị Cán bộ, Viên chức, Người lao động năm học 2026 – 2027.", True, "• Tuần 2 (Từ 07/9 – 13/9/2026): ")
    add_p("Thành lập các đội tuyển HSG và phân công giáo viên phụ trách ôn luyện; ra mắt và tổ chức sinh hoạt các Câu lạc bộ học sinh; tổ chức chuyên đề tuyên truyền ATGT và phòng chống tác hại thuốc lá điện tử.", True, "• Tuần 3 (Từ 14/9 – 20/9/2026): ")
    add_p("Kiểm tra nội bộ công tác chuyên môn, kiểm tra nề nếp dạy học và hồ sơ sổ sách đầu năm; chỉ đạo Đại hội Chi đoàn các lớp; họp giao ban Lãnh đạo trường và Tổ trưởng chuyên môn đánh giá công tác tháng 9, triển khai nhiệm vụ tháng 10/2026.", True, "• Tuần 4 (Từ 21/9 – 30/9/2026): ")

    # ==================== PHẦN KÝ TÊN VÀ NƠI NHẬN ====================
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False

    # Cột 1: Nơi nhận (Rộng ~8.0cm)
    c_nn = sig_tbl.cell(0, 0)
    c_nn.width = Cm(8.0)
    set_cell_margins(c_nn, top=0, bottom=0, left=0, right=100)

    p_nn_h = c_nn.paragraphs[0]
    p_nn_h.paragraph_format.space_before = Pt(0)
    p_nn_h.paragraph_format.space_after = Pt(2)
    r_nn_h = p_nn_h.add_run("Nơi nhận:")
    r_nn_h.font.name = 'Times New Roman'
    r_nn_h.font.size = Pt(12)
    r_nn_h.bold = True
    r_nn_h.italic = True

    recipients = [
        "- Sở GD&ĐT Đắk Lắk (để b/c);",
        "- Ban Giám hiệu (để chỉ đạo);",
        "- Các Tổ CM, Tổ Văn phòng (để t/h);",
        "- Đoàn TN, Công đoàn (để p/h);",
        "- Website nhà trường;",
        "- Lưu: VT, CM."
    ]

    for rc in recipients:
        p_r = c_nn.add_paragraph()
        p_r.paragraph_format.space_before = Pt(0)
        p_r.paragraph_format.space_after = Pt(1)
        p_r.paragraph_format.line_spacing = 1.0
        r_txt = p_r.add_run(rc)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(11)

    # Cột 2: Chức vụ & Chữ ký (Rộng ~8.5cm)
    c_sig = sig_tbl.cell(0, 1)
    c_sig.width = Cm(8.5)
    set_cell_margins(c_sig, top=0, bottom=0, left=100, right=0)

    p_pos = c_sig.paragraphs[0]
    p_pos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pos.paragraph_format.space_before = Pt(0)
    p_pos.paragraph_format.space_after = Pt(2)
    r_pos = p_pos.add_run("HIỆU TRƯỞNG")
    r_pos.font.name = 'Times New Roman'
    r_pos.font.size = Pt(13)
    r_pos.bold = True

    p_sign_note = c_sig.add_paragraph()
    p_sign_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sign_note.paragraph_format.space_before = Pt(0)
    p_sign_note.paragraph_format.space_after = Pt(40) # Khoảng trống ký tên ~4 dòng
    r_sn = p_sign_note.add_run("(Ký, đóng dấu và ghi rõ họ tên)")
    r_sn.font.name = 'Times New Roman'
    r_sn.font.size = Pt(11)
    r_sn.italic = True

    p_name = c_sig.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(0)
    r_nm = p_name.add_run("Lê Thị Thảo")
    r_nm.font.name = 'Times New Roman'
    r_nm.font.size = Pt(13)
    r_nm.bold = True

    output_filename = "Ke_Hoach_Cong_Tac_Thang_9_2026_THPT_CaoBaQuat.docx"
    doc.save(output_filename)
    print(f"SUCCESS: Generated {output_filename} ({os.path.getsize(output_filename)} bytes)")

if __name__ == "__main__":
    create_plan_docx()
